"""Construye la primera presentación (máximo 3 páginas, Arial 12) en Word y PDF.

Uso (desde esta carpeta):
    python3 construir_docx.py

Requiere python-docx y LibreOffice (soffice) para exportar a PDF. Escribe
../Primera_Presentacion_DW_Formula1.docx y ../Primera_Presentacion_DW_Formula1.pdf.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
SALIDA_DOCX = BASE.parent / "Primera_Presentacion_DW_Formula1.docx"

doc = Document()

# Todo en Arial 12, negro, interlineado sencillo
for nombre in ("Normal", "List Bullet"):
    estilo = doc.styles[nombre]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(12)
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = estilo.element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), "Arial")
    estilo.paragraph_format.space_before = Pt(0)
    estilo.paragraph_format.space_after = Pt(3)
    estilo.paragraph_format.line_spacing = 1.0

seccion = doc.sections[0]
seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
seccion.left_margin = seccion.right_margin = Cm(2.5)
seccion.top_margin = seccion.bottom_margin = Cm(2.5)


def _runs(parrafo, texto):
    """Agrega texto al párrafo; lo encerrado entre ** va en negrita."""
    for i, trozo in enumerate(texto.split("**")):
        if trozo:
            parrafo.add_run(trozo).bold = bool(i % 2)


def p(texto, centrado=False):
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY
    _runs(parrafo, texto)
    return parrafo


def titulo(texto):
    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.space_before = Pt(6)
    parrafo.paragraph_format.keep_with_next = True
    parrafo.add_run(texto).bold = True


def vinetas(elementos):
    for e in elementos:
        parrafo = doc.add_paragraph(style="List Bullet")
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parrafo.paragraph_format.space_after = Pt(1)
        _runs(parrafo, e)


def tabla(encabezados, filas, anchos_cm):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.autofit = False
    for fila_datos in [encabezados] + filas:
        celdas = t.rows[0].cells if fila_datos is encabezados else t.add_row().cells
        for j, valor in enumerate(fila_datos):
            par = celdas[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            if fila_datos is encabezados:
                par.add_run(valor).bold = True
            else:
                _runs(par, valor)
    for j, columna in enumerate(t.columns):
        columna.width = Cm(anchos_cm[j])
        for celda in columna.cells:
            celda.width = Cm(anchos_cm[j])
    for fila in t.rows:
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        fila._tr.get_or_add_trPr().append(cs)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------------------
# Encabezado
# ---------------------------------------------------------------------------
p("**[Universidad] – Base de Datos III – Proyecto: Implementación de un Data Warehouse**", centrado=True)
p("**Primera Presentación (puntos 1 a 5): Data Warehouse del Campeonato Mundial de Fórmula 1**", centrado=True)
p("Integrantes: Jhonatan Moisés Villca, [Integrante 2], [Integrante 3] – Docente: Rodnie Montaño Aguilera", centrado=True)

# ---------------------------------------------------------------------------
# 1. Tema
# ---------------------------------------------------------------------------
titulo("1. Tema del Proyecto")
p("**Tema:** análisis del rendimiento competitivo de pilotos, escuderías y circuitos en el Campeonato Mundial de Fórmula 1 (1950–2024). **Conjunto de datos:** «Formula 1 World Championship (1950–2024)» de Kaggle (Rohan Rao, datos de Ergast API): 14 archivos CSV relacionados por claves con ≈ 1 125 carreras, ≈ 26 700 resultados, ≈ 860 pilotos, ≈ 210 escuderías, 77 circuitos, 139 estados de finalización, posiciones del campeonato, clasificación, vueltas (≈ 589 000) y paradas en boxes.")
p("**Suficiencia de los datos:** contiene eventos medibles para la tabla de hechos (posición, puntos, vueltas y tiempos de cada resultado), entidades descriptivas con jerarquías para las dimensiones (piloto, escudería, circuito → país, fecha → temporada → década) y 75 años de historia. Además requiere limpieza (nulos escritos como «\\N», tiempos en texto, varios sistemas de puntos), lo que justifica el proceso ETL.")
p("**Contexto y necesidad de análisis:** cada temporada tiene Grandes Premios en distintos circuitos; los pilotos corren para una escudería y suman puntos según su posición final. Los datos están normalizados para registrar resultados, no para analizarlos en el tiempo. La necesidad es **evaluar de forma comparable entre épocas quién gana, con qué frecuencia y con qué confiabilidad**, para equipos, analistas, medios y patrocinadores. No es un tema de ventas, inventario ni similares: combina métricas ordinales, aditivas y semiaditivas, cambios de reglamento y varios procesos de negocio.")

# ---------------------------------------------------------------------------
# 2. Objetivos
# ---------------------------------------------------------------------------
titulo("2. Objetivos del Data Warehouse")
p("**Objetivo general:** implementar un Data Warehouse con un DataMart de Rendimiento en Carrera que integre los resultados históricos de la Fórmula 1 para analizar el rendimiento de pilotos, escuderías y circuitos mediante indicadores, consultas, cubos OLAP, reportes y cuadros de mando.")
p("**Objetivos específicos** (medibles con el indicador indicado):")
tabla(["Cód.", "Objetivo específico", "Indicador"], [
    ["OE1", "Medir la conversión de la pole en victoria por temporada y circuito.", "% victorias desde la pole"],
    ["OE2", "Determinar el dominio de las escuderías por temporada y década.", "% de victorias y cuota de puntos"],
    ["OE3", "Cuantificar la confiabilidad de las escuderías y las causas de abandono.", "Tasa de abandono"],
    ["OE4", "Comparar pilotos de distintas épocas con puntos normalizados.", "Puntos por carrera (sistema actual)"],
    ["OE5", "Caracterizar los circuitos por adelantamientos y abandonos.", "Prom. de posiciones ganadas"],
    ["OE6", "Evaluar la evolución del ritmo por circuito desde 2004.", "Velocidad media de vuelta rápida"],
], [1.4, 9.6, 5.0])
p("Estos objetivos definen el hecho (resultado de carrera), las dimensiones (tiempo, carrera, circuito, piloto, escudería y estado), las métricas, la granularidad y los reportes y cuadros de mando de la segunda presentación.")

# ---------------------------------------------------------------------------
# 3. Modelado conceptual
# ---------------------------------------------------------------------------
titulo("3. Modelado Conceptual del DataMart")
p("**Tres DataMarts posibles:**")
vinetas([
    "**DM1 Rendimiento en Carrera:** proceso resultado de carrera; cobertura completa 1950–2024.",
    "**DM2 Estrategia en Pista:** vueltas y paradas en boxes; solo desde 1996 y 2011.",
    "**DM3 Clasificación y Campeonato:** sesiones de clasificación y posiciones del campeonato.",
])
p("**DataMart seleccionado: DM1**, porque el resultado de carrera es el proceso central (de él salen los puntos y los campeonatos), tiene la historia completa, responde los seis objetivos y sus dimensiones se reutilizan en DM2 y DM3.")
p("**Modelo conceptual:** proceso de negocio «disputa de un Gran Premio»; hecho **Resultado de carrera**; dimensiones con sus jerarquías: Tiempo (década → año → mes → fecha), Carrera (temporada → Gran Premio), Circuito (continente → país → circuito), Piloto (nacionalidad → piloto), Escudería (nacionalidad → escudería) y Estado (categoría → estado).")
p("**Granularidad:** cada registro de la tabla de hechos es **el resultado de un piloto, con un auto de una escudería, en un Gran Premio** (una fila de results.csv); ≈ 26 700 registros.")

# ---------------------------------------------------------------------------
# 4. Modelado lógico
# ---------------------------------------------------------------------------
titulo("4. Modelado Lógico del DataMart")
p("Esquema en estrella. Cada dimensión se relaciona 1:N con la tabla de hechos mediante su clave sustituta (PK) y la clave foránea (FK) correspondiente:")
tabla(["Tabla", "Atributos"], [
    ["fact_resultado_carrera", "PK sk_resultado; FK sk_tiempo, sk_carrera, sk_circuito, sk_piloto, sk_escuderia, sk_estado; id_resultado_origen, codigo_posicion; métricas: puntos, puntos_sistema_actual, posicion_salida, posicion_final, posiciones_ganadas, vueltas_completadas, tiempo_vuelta_rapida_ms, velocidad_vuelta_rapida_kmh, edad_piloto, puntos_acumulados_temporada; indicadores 0/1: es_victoria, es_podio, es_pole, es_abandono, es_finalizado"],
    ["dim_tiempo", "PK sk_tiempo (AAAAMMDD); fecha, mes, trimestre, anio, decada"],
    ["dim_carrera", "PK sk_carrera; id_carrera_origen, temporada, ronda, nombre_gran_premio, sistema_puntos"],
    ["dim_circuito", "PK sk_circuito; id_circuito_origen, nombre_circuito, localidad, pais, continente"],
    ["dim_piloto", "PK sk_piloto; id_piloto_origen, nombre_completo, fecha_nacimiento, nacionalidad"],
    ["dim_escuderia", "PK sk_escuderia; id_escuderia_origen, nombre_escuderia, nacionalidad"],
    ["dim_estado", "PK sk_estado; id_estado_origen, descripcion_estado, categoria_estado"],
], [5.3, 10.7])
vinetas([
    "**Claves:** las claves sustitutas (sk_*) son enteros generados por el DBMS; las claves naturales del dataset (id_*_origen) son únicas y el ETL las usa para buscar la clave sustituta.",
    "**Métricas aditivas:** puntos, vueltas, posiciones ganadas e indicadores 0/1 (su suma cuenta victorias, podios o abandonos).",
    "**Métrica semiaditiva:** puntos acumulados de la temporada; se suma entre pilotos de una ronda, pero en el tiempo se toma el último valor.",
    "**Métricas no aditivas:** posiciones, tiempos, velocidad, edad y porcentajes.",
])
p("**Verificación:** OE1 usa es_pole y es_victoria por carrera y circuito; OE2, es_victoria y puntos por escudería y década; OE3, es_abandono por escudería y estado; OE4, puntos_sistema_actual por piloto; OE5, posiciones_ganadas por circuito; OE6, la vuelta rápida por circuito y temporada. Todos los objetivos se pueden responder con el modelo.")

# ---------------------------------------------------------------------------
# 5. Modelado físico
# ---------------------------------------------------------------------------
titulo("5. Modelado Físico del DataMart")
p("Se implementó en **PostgreSQL 16** (base dw_formula1, esquema dw_f1) con scripts SQL que crean las seis dimensiones y la tabla de hechos, los índices y la dimensión Tiempo (1950–2030).")
vinetas([
    "**Tipos de datos:** claves INTEGER/BIGINT IDENTITY; posiciones e indicadores SMALLINT; puntos NUMERIC(5,2) para admitir medios puntos; tiempos en milisegundos; fechas DATE.",
    "**Restricciones e integridad referencial:** PK en todas las tablas, UNIQUE en las claves naturales, seis FK obligatorias desde los hechos y reglas CHECK (posiciones > 0, puntos ≥ 0, es_victoria = 1 solo si posicion_final = 1). Se probó que el DBMS rechaza hechos con dimensiones inexistentes, resultados duplicados y el borrado de dimensiones con hechos.",
    "**Índices y orientación al análisis y al ETL:** índices en cada FK, en (piloto, tiempo), (escudería, tiempo) y en los atributos de las jerarquías; la clave natural única permite recargar sin duplicar y el miembro «Desconocido» (-1) evita perder hechos en la carga.",
])

doc.save(SALIDA_DOCX)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(SALIDA_DOCX.parent), str(SALIDA_DOCX)],
               check=True, capture_output=True)
print("Generados:", SALIDA_DOCX.name, "y", SALIDA_DOCX.with_suffix(".pdf").name)
