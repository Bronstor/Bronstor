"""Construye la primera presentación en Word y PDF (texto en Arial 12).

Uso (desde esta carpeta):
    NODE_PATH=$(npm root -g) python3 construir_docx.py

Requiere python-docx, LibreOffice (soffice) para exportar a PDF y Node.js con
Playwright para dibujar los diagramas. Escribe
../Primera_Presentacion_DW_Vuelos.docx y ../Primera_Presentacion_DW_Vuelos.pdf.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import diagramas

BASE = Path(__file__).resolve().parent
SALIDA_DOCX = BASE.parent / "Primera_Presentacion_DW_Vuelos.docx"

doc = Document()

# Todo en Arial 12, negro, interlineado sencillo
for nombre in ("Normal", "List Bullet"):
    estilo = doc.styles[nombre]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(12)
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = estilo.element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), "Arial")
    estilo.paragraph_format.space_before = Pt(0)
    estilo.paragraph_format.space_after = Pt(3)
    estilo.paragraph_format.line_spacing = 1.0

seccion = doc.sections[0]
seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
seccion.left_margin = seccion.right_margin = Cm(2.5)
seccion.top_margin = seccion.bottom_margin = Cm(2.5)


def _runs(parrafo, texto):
    """Agrega texto al párrafo; lo encerrado entre ** va en negrita."""
    for i, trozo in enumerate(texto.split("**")):
        if trozo:
            parrafo.add_run(trozo).bold = bool(i % 2)


def p(texto, centrado=False):
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY
    _runs(parrafo, texto)
    return parrafo


def titulo(texto):
    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.space_before = Pt(6)
    parrafo.paragraph_format.keep_with_next = True
    parrafo.add_run(texto).bold = True


def vinetas(elementos):
    for e in elementos:
        parrafo = doc.add_paragraph(style="List Bullet")
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parrafo.paragraph_format.space_after = Pt(1)
        _runs(parrafo, e)


def tabla(encabezados, filas, anchos_cm):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.autofit = False
    for fila_datos in [encabezados] + filas:
        celdas = t.rows[0].cells if fila_datos is encabezados else t.add_row().cells
        for j, valor in enumerate(fila_datos):
            par = celdas[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            if fila_datos is encabezados:
                par.add_run(valor).bold = True
            else:
                _runs(par, valor)
    for j, columna in enumerate(t.columns):
        columna.width = Cm(anchos_cm[j])
        for celda in columna.cells:
            celda.width = Cm(anchos_cm[j])
    for fila in t.rows:
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        fila._tr.get_or_add_trPr().append(cs)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figura(ruta_png, ancho, leyenda):
    imagen = doc.add_paragraph()
    imagen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imagen.paragraph_format.keep_with_next = True
    imagen.add_run().add_picture(str(ruta_png), width=ancho)
    p(leyenda, centrado=True)


def nueva_seccion(horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if horizontal else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Cm(29.7), Cm(21)) if horizontal else (Cm(21), Cm(29.7))
    s.top_margin = s.bottom_margin = Cm(2 if horizontal else 2.5)
    s.left_margin = s.right_margin = Cm(1.5 if horizontal else 2.5)


# ---------------------------------------------------------------------------
# Encabezado
# ---------------------------------------------------------------------------
p("**[Universidad] – Base de Datos III – Proyecto: Implementación de un Data Warehouse**", centrado=True)
p("**Primera Presentación (puntos 1 a 5): Data Warehouse de puntualidad de vuelos en EE. UU. (2015)**", centrado=True)
p("Integrantes: Jhonatan Moisés Villca, [Integrante 2], [Integrante 3] – Docente: Rodnie Montaño Aguilera", centrado=True)

# ---------------------------------------------------------------------------
# 1. Tema
# ---------------------------------------------------------------------------
titulo("1. Tema del Proyecto")
p("**Tema:** análisis de la puntualidad, los retrasos y las cancelaciones de los vuelos comerciales nacionales de Estados Unidos en 2015. **Conjunto de datos:** «2015 Flight Delays and Cancellations» de Kaggle, publicado por el Departamento de Transporte de EE. UU. con datos de la Oficina de Estadísticas de Transporte (BTS). Tiene tres archivos CSV: flights.csv (≈ 5,8 millones de vuelos con fecha, aerolínea, número de vuelo, aeropuertos de origen y destino, horarios programados y reales, retrasos, tiempos de rodaje y de vuelo, distancia, cancelación, desvío y minutos de retraso por causa), airlines.csv (14 aerolíneas) y airports.csv (322 aeropuertos con ciudad, estado y coordenadas).")
p("**Enlace del conjunto de datos:** https://www.kaggle.com/datasets/usdot/flight-delays")
p("**Suficiencia de los datos:** cada vuelo es un evento medible con métricas numéricas (minutos de retraso, tiempos, distancia) y con descriptores que forman dimensiones con jerarquías: año → trimestre → mes → fecha, franja horaria → hora y región → estado → ciudad → aeropuerto. El volumen y los problemas de calidad (horas en formato hhmm, valores vacíos en vuelos cancelados o desviados, aeropuertos con códigos numéricos en lugar del código IATA en octubre) justifican un Data Warehouse y un proceso ETL.")
p("**Contexto y necesidad de análisis:** según la BTS, un vuelo está retrasado si llega 15 minutos o más después de lo programado, y las aerolíneas informan la causa (la propia aerolínea, el clima, el sistema aéreo nacional, la seguridad o la llegada tardía del avión). Los vuelos están en un único archivo plano (flights.csv), útil para registrarlos pero no para analizarlos. La necesidad es **saber dónde, cuándo y por qué se producen los retrasos y las cancelaciones** (por aerolínea, aeropuerto, ruta, horario y causa), para que aerolíneas, aeropuertos y pasajeros tomen mejores decisiones. No es un tema de ventas ni de inventario: es el análisis del desempeño operativo de un sistema de transporte.")

# ---------------------------------------------------------------------------
# 2. Objetivos
# ---------------------------------------------------------------------------
titulo("2. Objetivos del Data Warehouse")
p("**Objetivo general:** implementar un Data Warehouse con un DataMart de Puntualidad de Vuelos que integre los vuelos de 2015 para analizar la puntualidad, los retrasos y las cancelaciones por aerolínea, aeropuerto, ruta y período, mediante indicadores, consultas, cubos OLAP, reportes y cuadros de mando.")
p("**Objetivos específicos** (cada uno se mide con el indicador indicado y se responde con una consulta SQL):")
tabla(["Cód.", "Objetivo específico", "Indicador (cómo se mide)"], [
    ["OE1", "Medir la puntualidad de cada aerolínea por mes.", "% puntuales = vuelos con menos de 15 min de retraso / vuelos completados"],
    ["OE2", "Cuantificar las cancelaciones y los desvíos por aerolínea y motivo.", "% cancelados y % desviados sobre los vuelos programados"],
    ["OE3", "Identificar qué causas generan más minutos de retraso.", "% de los minutos de retraso de cada causa"],
    ["OE4", "Determinar los aeropuertos y rutas con mayor retraso.", "Promedio de minutos de retraso de llegada por ruta"],
    ["OE5", "Analizar los retrasos según la franja horaria y el día de la semana.", "Promedio de minutos de retraso de salida por franja y día"],
    ["OE6", "Evaluar la eficiencia en tierra de cada aeropuerto.", "Promedio de minutos de rodaje antes del despegue"],
], [1.4, 8.2, 6.4])
p("Estos objetivos son la base del diseño: definen el hecho (vuelo) y su granularidad (un vuelo), las dimensiones (fecha, hora, aerolínea, aeropuerto y motivo de cancelación) y las métricas (minutos de retraso, tiempos, distancia e indicadores). En la segunda presentación cada objetivo tendrá su reporte (puntualidad por aerolínea, cancelaciones por motivo, causas de retraso, rutas y aeropuertos con más retraso, y retrasos por horario) y los indicadores se reunirán en un cuadro de mando de puntualidad con filtros por aerolínea, aeropuerto y mes.")

# ---------------------------------------------------------------------------
# 3. Modelado conceptual
# ---------------------------------------------------------------------------
titulo("3. Modelado Conceptual del DataMart")
p("**Tres DataMarts posibles:**")
vinetas([
    "**DM1 Puntualidad de Vuelos:** cada vuelo con sus retrasos, cancelación y desvío.",
    "**DM2 Causas de Retraso:** minutos de retraso de cada vuelo repartidos por causa.",
    "**DM3 Operaciones Aeroportuarias:** resumen diario de salidas, llegadas y cancelaciones de cada aeropuerto.",
])
p("**DataMart seleccionado: DM1**, porque tiene el detalle de cada vuelo (los otros dos se pueden obtener de él), responde los seis objetivos y sus dimensiones se reutilizan en DM2 y DM3.")
p("**Modelo conceptual:** proceso de negocio «operación de un vuelo programado»; hecho **Vuelo**; dimensiones con sus jerarquías: Fecha (año → trimestre → mes → fecha, y día de la semana), Hora de salida (franja horaria → hora), Aerolínea, Aeropuerto (región → estado → ciudad → aeropuerto), que se usa dos veces, como origen y como destino, y Motivo de cancelación (Figura 1).")
p("**Granularidad:** cada registro de la tabla de hechos es **un vuelo programado de una aerolínea en una fecha** (una fila de flights.csv), identificado por la fecha, la aerolínea, el número de vuelo, el aeropuerto de origen y la hora programada de salida; ≈ 5,8 millones de registros.")
figura(diagramas.exportar_png("conceptual"), Cm(15.5), "**Figura 1.** Modelo conceptual del DataMart Puntualidad de Vuelos: hecho, dimensiones y jerarquías.")

# ---------------------------------------------------------------------------
# 4. Modelado lógico
# ---------------------------------------------------------------------------
titulo("4. Modelado Lógico del DataMart")
p("Esquema en estrella (Figura 2). Cada dimensión se relaciona 1:N con la tabla de hechos mediante su clave sustituta (PK) y la clave foránea (FK) correspondiente:")
tabla(["Tabla", "Atributos"], [
    ["fact_vuelo", "PK sk_vuelo; FK sk_fecha, sk_hora_salida, sk_aerolinea, sk_aeropuerto_origen, sk_aeropuerto_destino, sk_motivo_cancelacion; dimensiones degeneradas: numero_vuelo, matricula_avion, salida_programada; métricas en minutos: retraso_salida_min, retraso_llegada_min, taxi_salida_min, taxi_llegada_min, tiempo_programado_min, tiempo_real_min, tiempo_aire_min, retraso_aerolinea_min, retraso_clima_min, retraso_sistema_aereo_min, retraso_seguridad_min, retraso_avion_tardio_min; distancia_millas; indicadores 0/1: es_cancelado, es_desviado, es_retrasado, es_puntual; fecha_carga (auditoría)"],
    ["dim_fecha", "PK sk_fecha (AAAAMMDD); fecha, dia, dia_semana, nombre_dia, es_fin_de_semana, mes, nombre_mes, trimestre, anio"],
    ["dim_hora", "PK sk_hora (0 a 23); hora_texto, franja_horaria"],
    ["dim_aerolinea", "PK sk_aerolinea; codigo_iata, nombre_aerolinea"],
    ["dim_aeropuerto", "PK sk_aeropuerto; codigo_iata, nombre_aeropuerto, ciudad, estado, region, pais, latitud, longitud"],
    ["dim_motivo_cancelacion", "PK sk_motivo_cancelacion; codigo_motivo (N, A, B, C, D), descripcion"],
], [5.3, 10.7])
p("**Estructura del dataset y su uso en el DataMart:** cada columna de los archivos CSV se carga en una dimensión, una métrica o un indicador de la tabla de hechos:")
tabla(["Columna del dataset", "Significado", "En el DataMart (uso)"], [
    ["YEAR\nMONTH\nDAY\nDAY_OF_WEEK", "Fecha del vuelo y día de la semana", "dim_fecha\n(dimensión)"],
    ["AIRLINE", "Código de la aerolínea", "dim_aerolinea\n(dimensión)"],
    ["ORIGIN_AIRPORT\nDESTINATION_AIRPORT", "Aeropuertos de origen y destino", "dim_aeropuerto, en dos roles\n(dimensión)"],
    ["SCHEDULED_DEPARTURE", "Hora programada de salida (hhmm)", "salida_programada (degenerada)\ndim_hora (dimensión)"],
    ["CANCELLATION_REASON", "Motivo: A, B, C o D", "dim_motivo_cancelacion\n(dimensión)"],
    ["FLIGHT_NUMBER\nTAIL_NUMBER", "Número de vuelo y matrícula del avión", "numero_vuelo\nmatricula_avion\n(dimensiones degeneradas)"],
    ["DEPARTURE_DELAY", "Minutos de retraso en la salida", "retraso_salida_min\n(métrica aditiva)"],
    ["ARRIVAL_DELAY", "Minutos de retraso en la llegada", "retraso_llegada_min (métrica aditiva)\nes_retrasado, es_puntual (indicadores)"],
    ["TAXI_OUT\nTAXI_IN", "Minutos de rodaje antes del despegue y después del aterrizaje", "taxi_salida_min\ntaxi_llegada_min\n(métricas aditivas)"],
    ["SCHEDULED_TIME\nELAPSED_TIME\nAIR_TIME", "Duración programada, real y en el aire", "tiempo_programado_min\ntiempo_real_min\ntiempo_aire_min\n(métricas aditivas)"],
    ["DISTANCE", "Distancia en millas", "distancia_millas\n(métrica aditiva)"],
    ["AIRLINE_DELAY\nWEATHER_DELAY\nAIR_SYSTEM_DELAY\nSECURITY_DELAY\nLATE_AIRCRAFT_DELAY", "Minutos de retraso por causa: aerolínea, clima, sistema aéreo, seguridad y avión tardío", "retraso_aerolinea_min\nretraso_clima_min\nretraso_sistema_aereo_min\nretraso_seguridad_min\nretraso_avion_tardio_min\n(métricas aditivas)"],
    ["CANCELLED\nDIVERTED", "1 si el vuelo fue cancelado o desviado", "es_cancelado\nes_desviado\n(indicadores 0/1)"],
    ["DEPARTURE_TIME\nWHEELS_OFF\nWHEELS_ON\nSCHEDULED_ARRIVAL\nARRIVAL_TIME", "Horas reales y programadas de cada etapa del vuelo", "No se cargan: lo necesario ya está en los retrasos y los tiempos"],
    ["airlines.csv: IATA_CODE, AIRLINE", "Código y nombre de la aerolínea", "dim_aerolinea\n(dimensión)"],
    ["airports.csv: IATA_CODE, AIRPORT, CITY, STATE, COUNTRY, LATITUDE, LONGITUDE", "Datos de cada aeropuerto", "dim_aeropuerto; la región se deriva del estado\n(dimensión)"],
], [6.2, 3.8, 6.0])
vinetas([
    "**Claves:** las claves sustitutas (sk_*) son enteros generados por el DBMS; los códigos IATA del dataset son claves naturales únicas que el ETL usa para buscar la clave sustituta. dim_aeropuerto se relaciona dos veces con los hechos (origen y destino).",
    "**Métricas aditivas:** minutos de retraso (total y por causa), tiempos, distancia e indicadores 0/1 (su suma cuenta vuelos cancelados, retrasados o puntuales).",
    "**Métricas semiaditivas:** no hay en este hecho, porque cada vuelo es un evento independiente y no un saldo acumulado.",
    "**Métricas no aditivas:** porcentajes de puntualidad y de cancelación, y promedios; se recalculan desde sus componentes.",
])
p("**Verificación:** OE1 usa es_puntual por aerolínea y mes; OE2, es_cancelado y es_desviado por aerolínea y motivo; OE3, los minutos por causa; OE4, retraso_llegada_min por aeropuertos de origen y destino; OE5, retraso_salida_min por franja horaria y día; OE6, taxi_salida_min por aeropuerto de origen. Todos los objetivos se pueden responder con el modelo.")

nueva_seccion(horizontal=True)
figura(diagramas.exportar_png("logico"), Cm(26.5), "**Figura 2.** Modelo lógico del DataMart Puntualidad de Vuelos (esquema en estrella).")
p("PK = clave primaria (sustituta); FK = clave foránea, con la tabla a la que apunta; NK = clave natural única; DD = dimensión degenerada; ‖──< = relación 1:N. dim_aeropuerto se relaciona dos veces con fact_vuelo (origen y destino).", centrado=True)
nueva_seccion(horizontal=False)

# ---------------------------------------------------------------------------
# 5. Modelado físico
# ---------------------------------------------------------------------------
titulo("5. Modelado Físico del DataMart")
p("Se implementó en **PostgreSQL 16** (base y esquema dw_vuelos) con scripts SQL que crean las cinco dimensiones y la tabla de hechos, los índices y los datos fijos (fechas 2014–2016, las 24 horas y los motivos de cancelación).")
vinetas([
    "**Tipos de datos:** claves INTEGER/BIGINT IDENTITY; minutos y distancia SMALLINT (los retrasos negativos indican adelanto); indicadores SMALLINT; fecha DATE y hora programada TIME; códigos VARCHAR(3).",
    "**Restricciones e integridad referencial:** PK en todas las tablas, UNIQUE en los códigos IATA y en la clave natural del vuelo, seis FK obligatorias (dos hacia dim_aeropuerto) y reglas CHECK (distancia > 0, un vuelo cancelado debe tener motivo, es_retrasado = 1 solo si llegó con 15 minutos o más de retraso). Se probó que el DBMS rechaza aerolíneas inexistentes, vuelos cancelados sin motivo, vuelos duplicados, distancias negativas y el borrado de aeropuertos con vuelos.",
    "**Índices y orientación al análisis y al ETL:** índices en cada FK, en (aerolínea, fecha), en la ruta (origen, destino) y un índice parcial de vuelos cancelados; la clave natural única permite recargar sin duplicar y el miembro «Desconocido» (-1) evita perder vuelos con códigos no encontrados.",
])
p("La Figura 3 muestra el modelo físico implementado, obtenido del catálogo de PostgreSQL: cada tabla con sus columnas, tipos de datos, claves y columnas obligatorias. Los scripts SQL de creación se entregan junto con el proyecto.")

# Figura del modelo físico (página horizontal)
nueva_seccion(horizontal=True)
figura(diagramas.exportar_png("fisico"), Cm(26.5), "**Figura 3.** Modelo físico del DataMart Puntualidad de Vuelos en PostgreSQL 16 (esquema dw_vuelos).")
p("PK = clave primaria; FK = clave foránea; UK = forma parte de una restricción UNIQUE; NN = NOT NULL; IDENTITY = valor generado por el DBMS; ‖──< = relación 1:N. dim_aeropuerto se relaciona dos veces con fact_vuelo (origen y destino).", centrado=True)

doc.save(SALIDA_DOCX)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(SALIDA_DOCX.parent), str(SALIDA_DOCX)],
               check=True, capture_output=True)
print("Generados:", SALIDA_DOCX.name, "y", SALIDA_DOCX.with_suffix(".pdf").name)
