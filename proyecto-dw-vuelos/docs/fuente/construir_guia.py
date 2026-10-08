"""Guía para exponer la primera presentación (Word y PDF, Arial 12).

Uso (desde esta carpeta):
    python3 construir_guia.py

Requiere python-docx y LibreOffice (soffice). Escribe
../Guia_Exposicion_DW_Vuelos.docx y ../Guia_Exposicion_DW_Vuelos.pdf.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
SALIDA_DOCX = BASE.parent / "Guia_Exposicion_DW_Vuelos.docx"

doc = Document()
for nombre in ("Normal", "List Bullet"):
    estilo = doc.styles[nombre]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(12)
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = estilo.element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), "Arial")
    estilo.paragraph_format.space_before = Pt(0)
    estilo.paragraph_format.space_after = Pt(4)
    estilo.paragraph_format.line_spacing = 1.0
seccion = doc.sections[0]
seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
seccion.left_margin = seccion.right_margin = Cm(2.5)
seccion.top_margin = seccion.bottom_margin = Cm(2.5)


def _runs(parrafo, texto, cursiva=False):
    for i, trozo in enumerate(texto.split("**")):
        if trozo:
            run = parrafo.add_run(trozo)
            run.bold = bool(i % 2)
            run.italic = cursiva


def p(texto, centrado=False):
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY
    _runs(parrafo, texto)


def titulo(texto):
    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.space_before = Pt(10)
    parrafo.paragraph_format.keep_with_next = True
    parrafo.add_run(texto).bold = True


def decir(texto):
    """Frase para decir en voz alta: con sangría y en cursiva."""
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    parrafo.paragraph_format.left_indent = Cm(1)
    _runs(parrafo, "«" + texto + "»", cursiva=True)


def vinetas(elementos):
    for e in elementos:
        parrafo = doc.add_paragraph(style="List Bullet")
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _runs(parrafo, e)


def tabla(encabezados, filas, anchos_cm):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.autofit = False
    for fila_datos in [encabezados] + filas:
        celdas = t.rows[0].cells if fila_datos is encabezados else t.add_row().cells
        for j, valor in enumerate(fila_datos):
            par = celdas[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            if fila_datos is encabezados:
                par.add_run(valor).bold = True
            else:
                _runs(par, valor)
    for j, columna in enumerate(t.columns):
        columna.width = Cm(anchos_cm[j])
        for celda in columna.cells:
            celda.width = Cm(anchos_cm[j])


p("**Guía para exponer la Primera Presentación**", centrado=True)
p("Data Warehouse de puntualidad de vuelos en EE. UU. (2015) – duración aproximada: 8 a 10 minutos", centrado=True)
p("Las frases en cursiva se pueden decir casi tal cual. Al mostrar los diagramas, recórrelos siempre del centro hacia afuera: primero el hecho y después cada dimensión.")

titulo("1. Tema del Proyecto (1 minuto)")
decir("Nuestro tema es el análisis de la puntualidad, los retrasos y las cancelaciones de los vuelos nacionales de Estados Unidos en 2015. Usamos el dataset 2015 Flight Delays and Cancellations de Kaggle, publicado por el Departamento de Transporte de EE. UU. Tiene tres archivos: uno con unos 5,8 millones de vuelos, otro con 14 aerolíneas y otro con 322 aeropuertos.")
decir("Elegimos este tema porque no es de ventas ni de inventario: es el análisis del funcionamiento de un sistema de transporte. La necesidad es saber dónde, cuándo y por qué se retrasan o cancelan los vuelos.")
p("**Dato clave:** un vuelo se considera retrasado si llega 15 minutos o más tarde; es la regla oficial del Departamento de Transporte.")

titulo("2. Objetivos del Data Warehouse (1 minuto)")
decir("Nuestro objetivo general es construir un Data Warehouse para analizar la puntualidad por aerolínea, aeropuerto, ruta y fecha.")
p("Lee rápido la tabla de los seis objetivos y remarca que cada uno se mide con un indicador. Por ejemplo:")
decir("La puntualidad se mide como los vuelos que llegaron con menos de 15 minutos de retraso, divididos entre los vuelos completados.")

titulo("3. Modelado Conceptual (2 minutos) – mostrar la Figura 1")
decir("Propusimos tres DataMarts: Puntualidad de Vuelos, Causas de Retraso y Operaciones Aeroportuarias. Elegimos Puntualidad de Vuelos porque tiene el detalle de cada vuelo, responde los seis objetivos y de él se pueden sacar los otros dos.")
p("Señalando la figura:")
decir("En el centro está el hecho Vuelo, con lo que medimos: minutos de retraso, tiempos, distancia y si el vuelo fue cancelado, desviado, retrasado o puntual. Alrededor están las dimensiones que lo describen: cuándo (fecha y hora), quién (aerolínea), desde y hacia dónde (aeropuerto de origen y de destino) y por qué se canceló (motivo).")
p("**Granularidad:**")
decir("Cada fila de la tabla de hechos es un vuelo programado de una aerolínea en una fecha.")

titulo("4. Modelado Lógico (2 minutos) – mostrar la Figura 2")
decir("Es un esquema en estrella: una tabla de hechos, fact_vuelo, y cinco dimensiones. Cada dimensión tiene una clave sustituta (PK), que es un número generado por la base de datos, y la tabla de hechos tiene una clave foránea (FK) hacia cada una. Todas las relaciones son de uno a muchos.")
p("Puntos que conviene explicar:")
vinetas([
    "**Aeropuerto aparece dos veces** (origen y destino), pero es una sola tabla. Esto se llama dimensión de rol.",
    "**Métricas aditivas:** los minutos, la distancia y los indicadores 0/1 se pueden sumar. Por ejemplo, sumar es_cancelado da el total de vuelos cancelados.",
    "**No hay métricas semiaditivas**, porque cada vuelo es un evento independiente, no un saldo que se acumula.",
    "**No aditivas:** los porcentajes y promedios. No se suman: se calculan de nuevo a partir de los totales.",
])

titulo("5. Modelado Físico (2 minutos) – mostrar la Figura 3")
decir("Lo implementamos en PostgreSQL 16. La figura sale directamente de la base de datos creada: cada columna con su tipo de dato, sus claves y si es obligatoria (NN significa NOT NULL).")
decir("La integridad referencial está garantizada por la base de datos. Lo probamos: rechaza un vuelo con una aerolínea que no existe, un vuelo cancelado sin motivo, un vuelo duplicado, una distancia negativa y el borrado de un aeropuerto que tiene vuelos.")
decir("Además pusimos índices para que las consultas sean rápidas, y la estructura está lista para cargar los datos con Apache Hop en la segunda presentación.")

titulo("Preguntas que pueden hacer")
tabla(["Pregunta", "Respuesta corta"], [
    ["¿Por qué no usar el código IATA como clave?", "Porque la clave sustituta es un número pequeño que hace más rápidas las uniones, no depende de la fuente y permite un registro «Desconocido» (-1). El código IATA se guarda como clave natural para buscar el registro en el ETL."],
    ["¿Qué es una dimensión degenerada?", "Un dato del vuelo que no necesita tabla propia, como el número de vuelo o la matrícula del avión. Se queda en la tabla de hechos."],
    ["¿Por qué los indicadores son SMALLINT y no BOOLEAN?", "Porque un número 0/1 se puede sumar para contar vuelos; un BOOLEAN no."],
    ["¿Qué pasa con el retraso de un vuelo cancelado?", "Queda vacío (NULL), porque el vuelo nunca llegó. Una regla CHECK de la base de datos lo controla."],
    ["¿Qué problemas tienen los datos?", "Las horas vienen como números (1530 = 15:30), hay valores vacíos en los vuelos cancelados y en octubre los aeropuertos aparecen con códigos numéricos en lugar del código IATA. Todo eso se corrige en el ETL."],
    ["¿Por qué el DataMart 1 y no los otros?", "Porque tiene el máximo detalle (cada vuelo) y responde todos los objetivos; los otros dos se pueden obtener de él."],
    ["¿Puede haber retrasos negativos?", "Sí: significa que el vuelo salió o llegó antes de la hora programada."],
    ["¿Por qué pidieron el enlace del dataset?", "Para comprobar que es público y de Kaggle, que tiene datos suficientes y que es el mismo que se cargará con Apache Hop en la segunda presentación."],
], [5.5, 10.5])

doc.save(SALIDA_DOCX)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(SALIDA_DOCX.parent), str(SALIDA_DOCX)],
               check=True, capture_output=True)
print("Generados:", SALIDA_DOCX.name, "y", SALIDA_DOCX.with_suffix(".pdf").name)
